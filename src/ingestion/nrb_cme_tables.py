"""Nepal liquidity driver series from archived NRB Current Macroeconomic Situation tables.

What: monthly series on the drivers of banking system liquidity in Nepal (workers'
remittances, overall balance of payments, gross foreign exchange reserves, government
treasury balance, government expenditure and revenue, NRB purchase of foreign currency,
NRB monetary operations, interbank volume) plus a few interest rates, collected from every
archived vintage of the "Current Macroeconomic (and Financial) Situation" statistical tables.

Source: Nepal Rastra Bank workbooks archived under data/raw/nepal/nrb_archive/macro_tables/
(see data/source_registry/manifests/nrb_wayback_manifest.csv for the original URLs).
Tables are located by their title text, never by sheet position, because sheet names and
layouts changed many times between 2006 and 2026. Two layouts are read:

* month row tables, where rows are mid month names (August ... July, or the Nepali month
  names) and columns are fiscal years (wide) or a fiscal year column runs down the left
  (the monthly annex series of newer vintages);
* snapshot tables (balance of payments, reserves, budgetary operations, structure of
  interest rates) that report a few fiscal year to date or mid month columns per vintage.

Period convention (as in src/features/nepal_interbank.py): a value labelled August of fiscal
year 2016/17 is the mid August value and is assigned to 2016-08; August to December belong
to the first calendar year of the fiscal year, January to July to the second.

Every observation is kept in the long file with its vintage (file name). The vintage
reference month is the last month of data the workbook covers, parsed from its title
("Based on Two Months data of 2083/84" means mid September 2026); data are public about two
months after it. Cells dated after the reference month are template cells and are dropped.
The monthly file keeps, for each series and period, the value from the most recent vintage
that reports it (latest reference month, then file name). Cumulative fiscal year to date
series get a derived monthly flow (suffix _flow) in the monthly file: the August flow equals
the August cumulative value, later months are the difference between consecutive cumulative
months of that table; a flow is left empty when the previous month is missing.

Series notes:
* workers_remittances, bop_overall: monthly annex (fiscal year 2016/17 onwards) where
  present, otherwise the balance of payments summary of the vintage. The older BPM5 table
  gives "Total, Group A through D" as the overall balance; the annex reports the BPM6 based
  overall balance, so the concept changes at 2016/17.
* gov_*: banking transaction (cash) basis from the budgetary operations table and the
  banking basis monthly annex; gov_*_fcgo: the FCGO based monthly annex (2019/20 onwards).
  A capital expenditure cell of zero next to a positive total is treated as not reported,
  because newer tables leave the breakdown at zero for the current year.
* nrb_liquidity_injection is the sum of repo, outright purchase, standing liquidity facility
  and overnight liquidity facility; nrb_liquidity_absorption the sum of reverse repo,
  outright sale, deposit collection, standing deposit facility and NRB bond, built in the
  monthly file from the latest value of each mo_* component (components that do not exist
  in a period simply do not enter). Tenors of one instrument are summed within a sheet; the
  long file holds the mo_* components only.
* interbank_rate, tbill_91d_wavg are kept as cross checks.

Units: Rs million throughout, except interest rates (percent). Annex reserves are reported in
Rs billion and are multiplied by 1000, except cells above 100000, which some vintages type in
Rs million.

Output:
  data/interim/nepal/nrb_cme_long.csv      every observation with its vintage
  data/interim/nepal/nrb_cme_monthly.csv   latest vintage wide table by period
  data/interim/nepal/nrb_cme_coverage.csv  coverage per series

Usage: python -m src.ingestion.nrb_cme_tables
"""

from __future__ import annotations

import datetime as dt
import re
import sys
import warnings
from collections import defaultdict
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
TABLES = ROOT / "data" / "raw" / "nepal" / "nrb_archive" / "macro_tables"
OUT_DIR = ROOT / "data" / "interim" / "nepal"
OUT_LONG = OUT_DIR / "nrb_cme_long.csv"
OUT_WIDE = OUT_DIR / "nrb_cme_monthly.csv"
OUT_COV = OUT_DIR / "nrb_cme_coverage.csv"
INTERBANK_REF = OUT_DIR / "nepal_interbank_monthly.csv"

SKIP_FILES = re.compile(r"QEB|Key_Financial_Indicators|Macroeconomic_Indicators", re.I)
HEAD_ROWS = 12

# ---------------------------------------------------------------- calendar helpers

EN_MONTHS = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6, "july": 7,
    "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
}
# Nepali month ending at mid of the Gregorian month given (Shrawan ends mid August).
NP_MONTHS = {
    "shrawan": 8, "srawan": 8, "saun": 8, "bhadra": 9, "bhadau": 9, "ashwin": 10, "aswin": 10,
    "asoj": 10, "kartik": 11, "mangsir": 12, "marga": 12, "poush": 1, "paush": 1, "push": 1,
    "magh": 2, "falgun": 3, "phalgun": 3, "fagun": 3, "chaitra": 4, "chait": 4, "baisakh": 5,
    "baishakh": 5, "jestha": 6, "jeth": 6, "asar": 7, "ashad": 7, "asadh": 7, "ashar": 7,
}
NUM_WORDS = {
    "one": 1, "first": 1, "two": 2, "second": 2, "three": 3, "third": 3, "four": 4,
    "five": 5, "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10, "eleven": 11,
    "twelve": 12, "annual": 12,
}
FY_RE = re.compile(r"((?:19|20)\d{2})\s*[/\-]\s*(\d{2,4})")


def month_of(text) -> int | None:
    """Gregorian month number of a mid month row label, or None."""
    if not isinstance(text, str):
        return None
    s = re.sub(r"[^a-z]", "", text.lower())
    if s.startswith("mid"):
        s = s[3:]
    if not s or len(s) > 11:
        return None
    if s == "ocotber":
        return 10
    for name, m in EN_MONTHS.items():
        if s == name or (len(s) >= 3 and name.startswith(s)) or s == "sept":
            return 9 if s == "sept" else m
    return NP_MONTHS.get(s)


def fy_start(text) -> int | None:
    """First Gregorian year of a fiscal year label such as 2016/17, 2024/2025R or 2072/73."""
    if not isinstance(text, str):
        return None
    for a, b in FY_RE.findall(text):
        y = int(a)
        nxt = int(b) if len(b) == 4 else int(b) + (y // 100) * 100
        if len(b) == 2 and nxt < y:
            nxt += 100
        if nxt == y + 1:
            return y - 57 if y >= 2035 else y  # Nepali (Bikram Sambat) fiscal year
    return None


def period_of(fy: int, month: int) -> str:
    year = fy if month >= 8 else fy + 1
    return f"{year}-{month:02d}"


def months_in_label(text) -> int | None:
    """Number of months of a fiscal year to date label: '8 Months', 'Six Months', 'Annual'."""
    if not isinstance(text, str):
        return None
    t = text.lower()
    if re.search(r"\bannual\b|twelve months|whole year", t):
        return 12
    m = re.search(r"(\d{1,2})\s*months?", t)
    if m and 1 <= int(m.group(1)) <= 12:
        return int(m.group(1))
    m = re.search(r"\b(one|first|two|second|three|four|five|six|seven|eight|nine|ten|eleven)\b\s*months?", t)
    if m:
        return NUM_WORDS[m.group(1)]
    return None


def vintage_of(name: str) -> tuple[int | None, int | None]:
    """(fiscal year start, months covered) from a workbook file name."""
    stem = re.sub(r"\.xlsx?$", "", name, flags=re.I)
    fy = None
    for a, b in re.findall(r"((?:19|20)\d{2})[._](\d{2})(?!\d)", stem):
        y = int(a)
        if (y + 1) % 100 == int(b):
            fy = y - 57 if y >= 2035 else y
    low = stem.lower().replace("_", " ").replace(".", " ")
    n = None
    if re.search(r"annual", low):
        n = 12
    else:
        m = re.search(r"(one|first|two|second|three|four|five|six|seven|eight|nine|ten|eleven|twelve)\s+(?:\d+\s+)?months?", low)
        if m:
            n = NUM_WORDS[m.group(1)]
        elif re.search(r"first\s+data|first\s+months", low):
            n = 1
    return fy, n


def ref_month(fy: int, n: int) -> str:
    month = (8 + n - 1 - 1) % 12 + 1
    return period_of(fy, month)


# ---------------------------------------------------------------- workbook reading

def _norm_cell(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        if v != v:  # NaN
            return None
        return float(v)
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return None  # date formatted cells hold no usable numbers in these tables
    s = str(v).strip()
    return s or None


def iter_sheets(path: Path):
    """Yield (sheet name, head text, loader) for every sheet; loader returns the full grid."""
    if path.suffix.lower() == ".xlsx":
        import openpyxl

        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            head = []
            for r in ws.iter_rows(max_row=HEAD_ROWS, values_only=True):
                head.extend(c for c in (_norm_cell(x) for x in r) if isinstance(c, str))

            def load(ws=ws):
                return [[_norm_cell(x) for x in r] for r in ws.iter_rows(values_only=True)]

            yield ws.title, _clean(" // ".join(head)), load
        wb.close()
    else:
        import xlrd

        wb = xlrd.open_workbook(path, on_demand=True, ignore_workbook_corruption=True)
        for name in wb.sheet_names():
            sh = wb.sheet_by_name(name)

            def cell(sh, i, j):
                c = sh.cell(i, j)
                if c.ctype == xlrd.XL_CELL_DATE:
                    return None
                return _norm_cell(c.value) if c.ctype in (1, 2) else None

            head = [cell(sh, i, j) for i in range(min(HEAD_ROWS, sh.nrows)) for j in range(sh.ncols)]
            head = [h for h in head if isinstance(h, str)]

            def load(sh=sh):
                return [[cell(sh, i, j) for j in range(sh.ncols)] for i in range(sh.nrows)]

            yield name, _clean(" // ".join(head)), load
        wb.release_resources()


def _clean(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace("\n", " ")).strip().lower()


def is_num(v) -> bool:
    return isinstance(v, float)


# ---------------------------------------------------------------- month row tables

def month_blocks(grid):
    """Split a sheet into month row blocks.

    Returns a list of dicts with keys: mcol (month label column), fycol (fiscal year column
    for long tables or None), rows (list of (row index, month)), labels (column -> label
    text from the header rows above), title (text of header cells left of the data).
    """
    ncols = max((len(r) for r in grid), default=0)
    hits = defaultdict(int)
    for r in grid:
        for j in range(min(5, len(r))):
            if month_of(r[j]):
                hits[j] += 1
    if not hits:
        return []
    mcol = max(hits, key=lambda j: (hits[j], -j))
    if hits[mcol] < 3:
        return []
    mrows = [(i, month_of(r[mcol])) for i, r in enumerate(grid) if len(r) > mcol and month_of(r[mcol])]
    blocks, cur, prev = [], [], -1
    for i, m in mrows:
        between = grid[prev + 1:i] if prev >= 0 else []
        header_like = any(
            any(isinstance(c, str) for c in r[mcol + 1:]) for r in between
        )
        if cur and header_like:
            blocks.append(cur)
            cur = []
        cur.append((i, m))
        prev = i
    if cur:
        blocks.append(cur)
    out, start = [], 0
    for b in blocks:
        first = b[0][0]
        header = grid[start:first]
        # fiscal year column of a long (annex) table: left of the month column
        fycol = None
        for j in range(mcol):
            if sum(1 for i, _ in b if j < len(grid[i]) and fy_start(grid[i][j])) >= 1:
                fycol = j
                break
        labels = defaultdict(list)
        title = []
        maxcol = max((j for r in header for j, v in enumerate(r) if j > mcol and isinstance(v, str)), default=-1)
        for r in header:
            last = None
            for j in range(min(ncols, maxcol + 1)):
                v = r[j] if j < len(r) else None
                if j <= mcol:
                    if isinstance(v, str):
                        title.append(v)
                    continue
                if isinstance(v, str):
                    last = v
                if last is not None:
                    labels[j].append(last)
        out.append({
            "mcol": mcol, "fycol": fycol, "rows": b,
            "labels": {j: _clean(" | ".join(v)) for j, v in labels.items()},
            "title": _block_title(title),
        })
        start = b[-1][0] + 1
    return out


def _block_title(parts):
    idx = [k for k, t in enumerate(parts) if re.match(r"\s*table\s*\d", t, re.I)]
    if idx:
        parts = parts[idx[-1]:]
    return _clean(" // ".join(parts))


def block_records(grid, block):
    """Yield (fy, month, column, label, value) for numeric cells of a month row block."""
    fy_row = None
    for i, m in block["rows"]:
        row = grid[i]
        if block["fycol"] is not None:
            f = fy_start(row[block["fycol"]]) if block["fycol"] < len(row) else None
            if f:
                fy_row = f
            # Annex tables restart in August; a missing label keeps the running year.
        for j in range(block["mcol"] + 1, len(row)):
            v = row[j]
            if not is_num(v):
                continue
            label = block["labels"].get(j, "")
            fy = fy_row if block["fycol"] is not None else fy_start(label)
            if fy is None:
                continue
            yield fy, m, j, label, v


# ---------------------------------------------------------------- observation store

class Store:
    def __init__(self):
        self.rows = []

    def add(self, series, fy, month, value, unit, basis, ctx, priority=0, period=None, key=""):
        period = period or period_of(fy, month)
        if ctx["ref"] and period > ctx["ref"]:
            return  # template cells beyond the data the vintage covers
        self.rows.append({
            "series": series,
            "period": period,
            "value": float(value),
            "unit": unit,
            "basis": basis,
            "source_file": ctx["file"],
            "sheet": ctx["sheet"],
            "vintage_reference_month": ctx["ref"],
            "_priority": priority,
            "_key": key,
        })


RS = "Rs million"
PCT = "percent"


def fy_window(ctx, fy) -> bool:
    """Accept only fiscal years the vintage can report (drops stale template columns)."""
    v = ctx["fy"]
    return v is None or v - 3 <= fy <= v


# ---------------------------------------------------------------- annex (long) tables

def parse_annex(grid, ctx, store, kind):
    for b in month_blocks(grid):
        if b["fycol"] is None:
            continue
        for fy, m, j, label, v in block_records(grid, b):
            if fy > (ctx["fy"] or 9999):
                continue
            if kind == "bop":
                if "remittance" in label and "%" not in label and "change" not in label:
                    store.add("workers_remittances", fy, m, v, RS, "cumulative_fy", ctx, 3)
                elif re.search(r"\bbop\b|overall", label):
                    store.add("bop_overall", fy, m, v, RS, "cumulative_fy", ctx, 3)
            elif kind == "reserves":
                if "rs" in label and "billion" in label and "gross foreign exchange reserve" in label:
                    # a few vintages type the newest month in Rs million inside the Rs billion
                    # column; such cells (above 100000) are used as they are
                    store.add("fx_reserves_npr", fy, m, v if v > 100000 else v * 1000.0, RS, "level", ctx, 3)
            elif kind == "treasury":
                if "treasury" in label:
                    store.add("treasury_position", fy, m, v, RS, "level", ctx, 3)
            elif kind in ("gbo_fcgo", "gbo_bank"):
                sfx = "_fcgo" if kind == "gbo_fcgo" else ""
                unit = RS + (", FCGO basis" if sfx else ", banking transaction basis")
                if re.search(r"total expenditure \| total expenditure$|total expenditure$", label):
                    store.add("gov_expenditure_total" + sfx, fy, m, v, unit, "cumulative_fy", ctx, 3)
                elif label.startswith("total expenditure") and label.endswith("capital"):
                    store.add("gov_expenditure_capital" + sfx, fy, m, v, unit, "cumulative_fy", ctx, 3)
                elif label.endswith("| total revenue") or label == "total revenue":
                    store.add("gov_revenue_total" + sfx, fy, m, v, unit, "cumulative_fy", ctx, 3)
                elif label == "expenditure":
                    store.add("gov_expenditure_total" + sfx, fy, m, v, unit, "cumulative_fy", ctx, 3)
                elif label == "revenue":
                    store.add("gov_revenue_total" + sfx, fy, m, v, unit, "cumulative_fy", ctx, 3)
            elif kind == "interest":
                if "interbank" in label:
                    store.add("interbank_rate", fy, m, v, PCT, "level", ctx, 3)
                elif "91" in label:
                    store.add("tbill_91d_wavg", fy, m, v, PCT, "level", ctx, 3)
                elif "deposit" in label:
                    store.add("deposit_rate_wavg", fy, m, v, PCT, "level", ctx, 3)
                elif "lending" in label:
                    store.add("lending_rate_wavg", fy, m, v, PCT, "level", ctx, 3)
                elif "base rate" in label:
                    store.add("base_rate", fy, m, v, PCT, "level", ctx, 3)


# ---------------------------------------------------------------- wide month tables

MONTH_KINDS = ("interbank_amount", "interbank_rate", "tbill91", "fx", "monetary_ops")


def mb_interbank_rate(grid, b, ctx, store, title):
    for fy, m, j, label, v in block_records(grid, b):
        if "other" in label or "overall" in label or fy > (ctx["fy"] or 9999):
            continue
        store.add("interbank_rate", fy, m, v, PCT, "level", ctx, 1)


def mb_interbank_amount(grid, b, ctx, store, title):
    labels = list(b["labels"].values())
    if any(re.search(r"\b[abc]\s*&\s*[abc]\b", l) for l in labels):
        return  # transactions by pair of bank classes, not the commercial bank series
    need_commercial = any("commercial" in l for l in labels)
    typed = any(re.search(r"amount|rate", l) for l in labels)
    for fy, m, j, label, v in block_records(grid, b):
        if fy > (ctx["fy"] or 9999) or re.search(r"other|overall", label):
            continue
        if need_commercial and "commercial" not in label:
            continue
        if typed:
            if "amount" in label:
                store.add("interbank_volume", fy, m, v, RS, "flow", ctx, 2)
            elif "rate" in label:
                store.add("interbank_rate", fy, m, v, PCT, "level", ctx, 2)
        elif "rate" not in title.split("transaction")[-1][:20]:
            store.add("interbank_volume", fy, m, v, RS, "flow", ctx, 1)


def mb_tbill91(grid, b, ctx, store, title):
    for fy, m, j, label, v in block_records(grid, b):
        if fy > (ctx["fy"] or 9999):
            continue
        if "91" in label or ("91" in title and not re.search(r"364|182|28", label)):
            store.add("tbill_91d_wavg", fy, m, v, PCT, "level", ctx, 1)


def mb_fx(grid, b, ctx, store, title):
    def rs_col(l):
        return re.fullmatch(r"n?rs\.?|npr", l.split("|")[-1].strip()) is not None

    labels = b["labels"]
    has_rs = any(rs_col(l) for l in labels.values())
    text = title + " " + ctx.get("head", "")
    usd_table = bool(re.search(r"us\$|usd", text)) and not re.search(r"\(rs|rs\.? in", text)
    for fy, m, j, label, v in block_records(grid, b):
        if not fy_window(ctx, fy):
            continue
        if re.search(r"\bic\b|ic purchase|indian|us\$ sale", label):
            continue
        if not re.search(r"\bnet\b", label):
            continue
        if has_rs:
            if not rs_col(label):
                continue
        elif usd_table or re.search(r"us\$|usd", label):
            continue
        store.add("nrb_fx_net_purchase", fy, m, v, RS, "flow", ctx, 1)


MO_CLASSES = [
    (r"reverse\s*repo", "reverse_repo", "absorption"),
    (r"outright\s*sale", "outright_sale", "absorption"),
    (r"deposit\s*(collection|auction)", "deposit_collection", "absorption"),
    (r"standing\s*deposit|\bsdf\b", "standing_deposit_facility", "absorption"),
    (r"nrb\s*bond", "nrb_bond", "absorption"),
    (r"outright\s*purchase", "outright_purchase", "injection"),
    (r"overnight\s*liquidity|\bolf\b", "overnight_liquidity_facility", "injection"),
    (r"standing\s*liquidity|\bslf\b", "standing_liquidity_facility", "injection"),
    (r"\brepo\b", "repo", "injection"),
]


def classify_mo(text):
    for pat, comp, side in MO_CLASSES:
        if re.search(pat, text):
            return comp, side
    return None, None


def mb_monetary_ops(grid, b, ctx, store, title):
    labels = b["labels"]
    has_total = any(re.search(r"\btotal\b", l) for l in labels.values())
    for fy, m, j, label, v in block_records(grid, b):
        if not fy_window(ctx, fy):
            continue
        if re.search(r"rate|%|interest|number|\bno\.|outstanding|maturity|\bdays? limit", label):
            continue
        if has_total and not re.search(r"\btotal\b", label):
            continue  # facility split by institution class: keep the total column only
        comp, side = classify_mo(label)
        if comp is None:
            comp, side = classify_mo(title)
        if comp is None:
            continue
        # sub instruments (tenors) of one facility are summed later through the key
        store.add(f"mo_{comp}", fy, m, v, RS, "flow", ctx, 1, key=label)


MB_FUNCS = {
    "interbank_amount": mb_interbank_amount,
    "interbank_rate": mb_interbank_rate,
    "tbill91": mb_tbill91,
    "fx": mb_fx,
    "monetary_ops": mb_monetary_ops,
}


def parse_transposed(grid, ctx, store, kind):
    """Fiscal years down the rows and mid months across the columns."""
    series = "interbank_rate" if kind == "interbank_rate" else "tbill_91d_wavg"
    for hi, r in enumerate(grid):
        mcols = {j: month_of(c) for j, c in enumerate(r) if month_of(c)}
        if len(mcols) < 10:
            continue
        for r2 in grid[hi + 1:]:
            fy = None
            for c in r2[:3]:
                fy = fy_start(c)
                if fy:
                    break
            if fy is None or fy > (ctx["fy"] or 9999):
                continue
            for j, m in mcols.items():
                if j < len(r2) and is_num(r2[j]):
                    store.add(series, fy, m, r2[j], PCT, "level", ctx, 1)
        return


def parse_month_tables(grid, ctx, store, head, sheet_kind):
    blocks = [b for b in month_blocks(grid) if b["fycol"] is None]
    if not blocks:
        if sheet_kind in ("interbank_rate", "tbill91"):
            parse_transposed(grid, ctx, store, sheet_kind)
        return
    for b in blocks:
        kind, title = sheet_kind, head
        if re.match(r"table\s*\d", b["title"]):
            own = [k for k in classify(b["title"], "") if k in MONTH_KINDS]
            if own:
                kind, title = own[0], b["title"]
            elif b["title"] != head[: len(b["title"])]:
                continue  # a different kind of table stacked in the same sheet
        MB_FUNCS[kind](grid, b, ctx, store, title)


# ---------------------------------------------------------------- snapshot tables

def header_columns(grid, first_data_row, start_col=0):
    """Column label text built from rows above the first data row (forward filled)."""
    ncols = max((len(r) for r in grid), default=0)
    labels = defaultdict(list)
    for r in grid[:first_data_row]:
        last = None
        for j in range(start_col, ncols):
            v = r[j] if j < len(r) else None
            if isinstance(v, str) or (is_num(v) and 1990 <= v <= 2100 and float(v).is_integer()):
                last = v if isinstance(v, str) else str(int(v))
            if last is not None:
                labels[j].append(last)
    return {j: _clean(" | ".join(v)) for j, v in labels.items()}


def row_label(row, upto=7):
    parts = [c for c in row[:upto] if isinstance(c, str)]
    return _clean(" ".join(parts))


def first_numeric_row(grid, min_cols=2):
    for i, r in enumerate(grid):
        nums = [c for c in r if is_num(c)]
        if len(nums) >= min_cols and row_label(r) and not all(c.is_integer() and 1990 <= c <= 2100 for c in nums):
            return i
    return None


def label_col_end(grid, i0):
    """Index of the last column holding row labels (data columns start after it)."""
    end = 0
    for r in grid[i0:]:
        for j, c in enumerate(r[:8]):
            if isinstance(c, str) and not FY_RE.search(c) and c.strip() not in ("-",):
                end = max(end, j)
    return end


def header_model(grid, i0, start_col):
    """Column headers of a snapshot table.

    A header cell spans to the right over empty cells only while some deeper header row has
    a label in that column (the usual merged cell layout). Returns, per column, the label
    text, the fiscal year (deepest own label first) and the number of fiscal year to date
    months named in a period label such as 'Eight Months' or 'Annual'.
    """
    ncols = max((len(r) for r in grid), default=0)
    hdr = grid[:i0]

    def cell(r, j):
        v = r[j] if j < len(r) else None
        if isinstance(v, str):
            return v
        if is_num(v) and v.is_integer() and 1990 <= v <= 2100:
            return str(int(v))
        return None

    own = [[cell(r, j) for j in range(ncols)] for r in hdr]
    below = [[any(own[k][j] for k in range(i + 1, len(hdr))) for j in range(ncols)] for i in range(len(hdr))]
    text, fy_own, fy_fill = defaultdict(list), {}, {}
    per_own, per_span, per_carry = {}, {}, {}
    for i, r in enumerate(own):
        last = last_per = last_fy = None
        for j in range(start_col, ncols):
            v = r[j]
            if v is not None:
                n, f = months_in_label(v), fy_start(v)
                text[j].append(v)
                last = v
                if n:
                    per_own[j] = last_per = n
                elif f and last_per:
                    per_carry[j] = last_per  # a bare year label inside a period span
                if f:
                    fy_own[j] = last_fy = f
                if not n and not f:
                    last_per = last_fy = None
            elif below[i][j]:
                if last is not None:
                    text[j].append(last)
                if last_per:
                    per_span[j] = last_per
                if last_fy:
                    fy_fill[j] = last_fy
            else:
                last = last_per = last_fy = None
    out = {}
    for j in range(start_col, ncols):
        out[j] = {"text": _clean(" | ".join(text.get(j, []))), "fy": fy_own.get(j) or fy_fill.get(j),
                  "months": per_own.get(j) or per_span.get(j) or per_carry.get(j)}
    return out


def fytd_columns(grid, ctx, i0, head):
    """Map data columns to (fy, months, label) for fiscal year to date snapshot tables."""
    lce = label_col_end(grid, i0)
    model = header_model(grid, i0, lce + 1)
    explicit = any(c["months"] for c in model.values())
    head_months = months_in_label(head) or ctx.get("months")
    cols = {}
    for j, c in model.items():
        lab, fy = c["text"], c["fy"]
        if not lab or re.search(r"change|percent|%|share|ratio|composition|growth", lab):
            continue
        if fy is None or not fy_window(ctx, fy):
            continue
        n = c["months"]
        if n is None:
            if explicit or re.search(r"credit|debit|budget|estimate|target|revised|projection", lab):
                continue
            n = head_months
        if n is None:
            continue
        if ctx.get("months") and n not in (ctx["months"], 12):
            continue  # stale label left over from an older template
        cols[j] = (fy, n, lab)
    return cols


def parse_bop_snapshot(grid, ctx, store, head):
    i0 = first_numeric_row(grid)
    if i0 is None:
        return
    lce = label_col_end(grid, i0)
    if any(re.search(r"credit|debit", c["text"]) for c in header_model(grid, i0, lce + 1).values()):
        return  # BPM6 credit, debit, net layout: covered by the monthly annex where present
    cols = fytd_columns(grid, ctx, i0, head)
    if not cols:
        return
    seen = set()
    for r in grid[i0:]:
        lab = row_label(r)
        if not lab:
            continue
        if re.search(r"workers'?\s*remittance", lab) and "remit" not in seen:
            series = "workers_remittances"
            seen.add("remit")
        elif re.search(r"total,?\s*groups?\s*a\s*through\s*d|overall\s*balance", lab) and "bop" not in seen:
            series = "bop_overall"
            seen.add("bop")
        else:
            continue
        for j, (fy, n, _) in cols.items():
            if j < len(r) and is_num(r[j]):
                store.add(series, fy, (8 + n - 2) % 12 + 1, r[j], RS, "cumulative_fy", ctx, 1)


def strip_numbering(lab: str) -> str:
    prev = None
    lab = lab.strip()
    while prev != lab:
        prev = lab
        lab = re.sub(r"^(\d+(\.\d+)*[\.\)]?\s+|\(?[ivx]+[\.\)]\s*|\(?[a-h][\.\)]\s*|[\-–]\s*)", "", lab).strip()
    return lab.strip(" :+*#")


def parse_gbo_snapshot(grid, ctx, store, head):
    i0 = first_numeric_row(grid)
    if i0 is None:
        return
    cols = fytd_columns(grid, ctx, i0, head)
    if not cols:
        return
    unit = f"{RS}, banking transaction basis"
    rows = [(strip_numbering(row_label(r)), r) for r in grid[i0:]]
    rows = [(c, r) for c, r in rows if c]

    def first(pat, start=0, exclude=None):
        for k in range(start, len(rows)):
            c = rows[k][0]
            if re.match(pat, c) and not (exclude and re.search(exclude, c)):
                return k
        return None

    k_total = first(r"total expenditure")
    k_actual = first(r"actual expend[iu]ture")
    k_budget = first(r"expenditure of budget")
    k_tot = next((k for k in (k_total, k_actual, k_budget) if k is not None), None)
    anchors = [k for k in (k_actual, k_budget, k_total) if k is not None]
    if k_actual is not None:
        anchors = [k_actual]
    k_cap = first(r"capital\b", min(anchors) + 1) if anchors else None
    k_res = first(r"(total\s+)?resources")
    k_rev = first(r"revenue\b", k_res + 1 if k_res is not None else 0, exclude=r"grant")
    tot = rows[k_tot][1] if k_tot is not None else []
    for k, series in ((k_tot, "gov_expenditure_total"), (k_cap, "gov_expenditure_capital"),
                      (k_rev, "gov_revenue_total")):
        if k is None:
            continue
        r = rows[k][1]
        for j, (fy, n, _) in cols.items():
            if j < len(r) and is_num(r[j]):
                if series == "gov_expenditure_capital" and r[j] == 0 and j < len(tot) and is_num(tot[j]) and tot[j] > 0:
                    continue  # breakdown not filled in for this column (only the total is)
                store.add(series, fy, (8 + n - 2) % 12 + 1, r[j], unit, "cumulative_fy", ctx, 1)


def parse_gbo_cash_balance(grid, ctx, store, head):
    """Cash balance of the government (treasury position) from the newer budgetary table."""
    i0 = first_numeric_row(grid)
    if i0 is None:
        return
    cols = fytd_columns(grid, ctx, i0, head)
    for r in grid[i0:]:
        lab = row_label(r)
        if re.search(r"cash balance of (the )?general government", lab):
            for j, (fy, n, _) in cols.items():
                if j < len(r) and is_num(r[j]):
                    store.add("treasury_position", fy, (8 + n - 2) % 12 + 1, r[j], RS, "level", ctx, 1)
            return


GREG_MONTH = re.compile(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?", re.I)
YEAR = re.compile(r"\b((?:19|20)\d{2})\b")


def gregorian_columns(grid, i0, start_col):
    """Map columns to Gregorian periods from header cells such as 'Mid Jul' and '2015'."""
    out = {}
    ncols = max((len(r) for r in grid), default=0)
    model = header_model(grid, i0, start_col)
    for j in range(start_col, ncols):
        if re.search(r"change|percent|%", model[j]["text"]) or FY_RE.search(model[j]["text"]):
            continue
        cells = []
        for r in grid[:i0]:
            v = r[j] if j < len(r) else None
            if isinstance(v, str):
                cells.append(v)
            elif is_num(v) and 1990 <= v <= 2100 and float(v).is_integer():
                cells.append(str(int(v)))
        text = " ".join(cells)
        if re.search(r"change|percent|%|to\s+mid", text, re.I):
            continue
        txt = re.sub(r"(?i)mid[\s\-]*", " ", text)
        mo = GREG_MONTH.findall(txt)
        yr = [y for y in YEAR.findall(txt) if not FY_RE.search(txt)] or YEAR.findall(txt)
        if len(mo) == 1 and len(yr) == 1:
            m = month_of(mo[0])
            if m:
                out[j] = f"{int(yr[0])}-{m:02d}"
    return out


def parse_reserves_snapshot(grid, ctx, store, head):
    i0 = first_numeric_row(grid)
    if i0 is None:
        return
    lce = label_col_end(grid, i0)
    cols = gregorian_columns(grid, i0, lce + 1)
    if not cols:
        return
    chosen = None
    for r in grid[i0:]:
        lab = row_label(r)
        if re.search(r"gross foreign exchange reserve", lab) and not re.search(r"import|month|share", lab):
            chosen = r
            break
    if chosen is None:
        for r in grid[i0:]:
            if re.match(r"total reserves?$", row_label(r)):
                chosen = r
                break
    if chosen is None:
        return
    for j, period in cols.items():
        if j < len(chosen) and is_num(chosen[j]):
            store.add("fx_reserves_npr", None, None, chosen[j], RS, "level", ctx, 1, period=period)


def parse_interest_structure(grid, ctx, store):
    """Structure of Interest Rates: periods across columns, rates down the rows."""
    wanted = [
        (r"weighted average interbank", "interbank_rate"),
        (r"weighted average deposit", "deposit_rate_wavg"),
        (r"weighted average lending", "lending_rate_wavg"),
        (r"base rate", "base_rate"),
        (r"t-?\s*bills?\*?\s*\(?\s*91", "tbill_91d_wavg"),
    ]
    # tables may be stacked; handle each header ('Year' row) separately
    starts = [i for i, r in enumerate(grid) if r and isinstance(r[0], str) and re.match(r"\s*year", r[0], re.I)]
    for k, s in enumerate(starts):
        end = starts[k + 1] if k + 1 < len(starts) else len(grid)
        sub = grid[s:end]
        i0 = None
        for i, r in enumerate(sub):
            if i > 0 and any(is_num(c) for c in r[1:]):
                i0 = i
                break
        if i0 is None:
            continue
        hdr = sub[: min(i0, 3)]
        cols = {}
        ncols = max(len(r) for r in sub)
        for j in range(1, ncols):
            cells = []
            for r in hdr:
                v = r[j] if j < len(r) else None
                if isinstance(v, str):
                    cells.append(v)
                elif is_num(v) and 1990 <= v <= 2100:
                    cells.append(str(int(v)))
            text = " ".join(cells)
            mo, yr = GREG_MONTH.findall(text), YEAR.findall(text)
            if len(mo) == 1 and len(yr) == 1 and month_of(mo[0]):
                cols[j] = f"{int(yr[0])}-{month_of(mo[0]):02d}"
        if not cols:
            continue
        done = set()
        for r in sub[i0:]:
            lab = row_label(r, 3)
            for pat, series in wanted:
                if series in done or not re.search(pat, lab):
                    continue
                done.add(series)
                for j, period in cols.items():
                    if j < len(r) and is_num(r[j]):
                        store.add(series, None, None, r[j], PCT, "level", ctx, 0, period=period)
                break


# ---------------------------------------------------------------- dispatcher

def classify(head: str, sheet: str) -> list[str]:
    h = head[:400]
    kinds = []
    annex = h.startswith("annex") or " annex " in h[:40] or "time series" in h[:80]
    if "balance of payments indicators" in h and "monthly" in h:
        kinds.append("annex_bop")
    elif "foreign exchange reserves" in h and "mid-month" in h and "annual" not in h:
        kinds.append("annex_reserves")
    elif "government budgetary operations (monthly series)" in h:
        kinds.append("annex_gbo_fcgo" if "fcgo" in h else "annex_gbo_bank")
        if re.search(r"treasury (balance|position)", h):
            kinds.append("annex_treasury")
    elif re.search(r"treasury (balance|position)", h):
        kinds.append("annex_treasury")
    elif re.search(r"interest rates", h) and "fiscal year" in h and "mid-month" in h:
        kinds.append("annex_interest")
    elif annex:
        return kinds
    elif re.search(r"weighted average( \| | // )?\s*inter-?\s*bank (transaction )?rate", h):
        kinds.append("interbank_rate")
    elif re.search(r"inter-?\s*bank transaction(?! rate)", h):
        kinds.append("interbank_amount")
    elif re.search(r"treasury bills? rate", h) and "91" in h:
        kinds.append("tbill91")
    elif re.search(r"purchase\s*/\s*sale of foreign currency|foreign exchange intervention", h):
        kinds.append("fx")
    elif re.search(r"monetary operation|outright (sale|purchase)|repo auction|standing liquidity facility|deposit collection|liquidity management", h) and "summary" not in h:
        kinds.append("monetary_ops")
    elif "structure of interest rate" in h:
        kinds.append("interest_structure")
    elif "balance of payments" in h and not re.search(r"us\$|usd|current month|mid-?\s*\w+ to mid|mid \w+ - mid|annual series", h):
        kinds.append("bop")
    elif re.search(r"gross foreign (exchange )?(holding|assets|reserves?)\b.*banking sector|foreign holding of the banking", h) \
            and not re.search(r"us\$|usd", h):
        kinds.append("reserves")
    elif re.search(r"government budgetary operation", h) and not re.search(r"annual series|growth|percent of gdp|monthly series", h):
        kinds.append("gbo")
    return kinds


def parse_file(path: Path, store: Store, report: list):
    fy, n = vintage_of(path.name)
    ctx = {"file": path.name, "fy": fy, "months": n, "ref": ref_month(fy, n) if fy and n else ""}
    found = defaultdict(list)
    try:
        for sheet, head, load in iter_sheets(path):
            kinds = classify(head, sheet)
            if not kinds:
                continue
            grid = load()
            c = dict(ctx, sheet=sheet, head=head)
            before = len(store.rows)
            for k in kinds:
                if k.startswith("annex_"):
                    parse_annex(grid, c, store, k[6:])
                elif k in MONTH_KINDS:
                    parse_month_tables(grid, c, store, head, k)
                elif k == "interest_structure":
                    parse_interest_structure(grid, c, store)
                elif k == "bop":
                    parse_bop_snapshot(grid, c, store, head)
                elif k == "reserves":
                    parse_reserves_snapshot(grid, c, store, head)
                elif k == "gbo":
                    parse_gbo_snapshot(grid, c, store, head)
                    parse_gbo_cash_balance(grid, c, store, head)
                found[k].append(f"{sheet}:{len(store.rows) - before}")
    except Exception as e:  # unreadable workbook: report and continue
        report.append({"file": path.name, "ref": ctx["ref"], "status": f"error {type(e).__name__}: {e}"})
        return
    report.append({"file": path.name, "ref": ctx["ref"], "status": "ok",
                   **{k: ";".join(v) for k, v in found.items()}})


# ---------------------------------------------------------------- assembly

CUMULATIVE = {"workers_remittances", "bop_overall", "gov_expenditure_total", "gov_expenditure_capital",
              "gov_revenue_total", "gov_expenditure_total_fcgo", "gov_expenditure_capital_fcgo",
              "gov_revenue_total_fcgo"}
INJ = ["repo", "outright_purchase", "standing_liquidity_facility", "overnight_liquidity_facility"]
ABS = ["reverse_repo", "outright_sale", "deposit_collection", "standing_deposit_facility", "nrb_bond"]


def per_vintage(long: pd.DataFrame) -> pd.DataFrame:
    """One value per series, period and vintage file (highest table priority, then first seen).

    When one sheet gives two different values for the same series and period (a shifted
    year header), neither is used for that vintage; both stay in the long file.
    """
    n = long.groupby(["series", "period", "source_file", "sheet"]).value.transform("nunique")
    conflicts = long[n > 1]
    if len(conflicts):
        print(f"conflicting values inside one sheet (left out of the monthly table): "
              f"{conflicts.groupby(['series', 'period', 'source_file', 'sheet']).ngroups} cases")
    long = long[n == 1]
    long = long.sort_values(["series", "period", "source_file", "_priority"], ascending=[True, True, True, False],
                            kind="stable")
    return long.drop_duplicates(["series", "period", "source_file"], keep="first")


def add_mo_totals(wide: pd.DataFrame, latest: pd.DataFrame):
    """Injection and absorption totals: sum of the latest value of each component."""
    rows = []
    for name, keys in (("nrb_liquidity_injection", INJ), ("nrb_liquidity_absorption", ABS)):
        comps = [f"mo_{k}" for k in keys if f"mo_{k}" in wide]
        if not comps:
            continue
        wide[name] = wide[comps].sum(axis=1, min_count=1)
        part = latest[latest.series.isin(comps)].assign(series=name)
        rows.append(part[part.period.isin(wide.index[wide[name].notna()])])
    return wide, pd.concat([latest, *rows], ignore_index=True)


def prev_month(p: str) -> str:
    y, m = int(p[:4]), int(p[5:])
    return f"{y}-{m - 1:02d}" if m > 1 else f"{y - 1}-12"


def add_flows(wide: pd.DataFrame, latest: pd.DataFrame):
    """Monthly flows of the cumulative fiscal year to date series (latest vintage values).

    The August flow equals the August cumulative value; later months are the difference
    from the previous month, which always lies in the same fiscal year (July closes it).
    A flow is left empty when the previous month is missing.
    """
    flows, rows = {}, []
    for s in sorted(c for c in wide.columns if c in CUMULATIVE):
        col = wide[s]
        out = {}
        for p, v in col.dropna().items():
            if p.endswith("-08"):
                out[p] = v
            else:
                prev = col.get(prev_month(p))
                if prev is not None and pd.notna(prev):
                    out[p] = v - prev
        flows[s + "_flow"] = pd.Series(out, dtype=float)
        src = latest[latest.series == s].set_index("period")
        for p in out:
            rows.append({"series": s + "_flow", "period": p, "source_file": src.at[p, "source_file"]})
    wide = pd.concat([wide, pd.DataFrame(flows)], axis=1).sort_index()
    wide = wide[sorted(wide.columns)]
    wide.index.name = "period"
    return wide, pd.concat([latest, pd.DataFrame(rows)], ignore_index=True)


def vintage_key(df: pd.DataFrame) -> pd.Series:
    # most recent vintage: latest reference month, then latest upload prefix and file name
    return df.vintage_reference_month.fillna("") + "|" + df.source_file


def build_wide(pv: pd.DataFrame) -> pd.DataFrame:
    pv = pv.assign(_vk=vintage_key(pv)).sort_values(["series", "period", "_vk"])
    latest = pv.drop_duplicates(["series", "period"], keep="last")
    wide = latest.pivot(index="period", columns="series", values="value").sort_index()
    wide.index.name = "period"
    return wide, latest


def coverage(latest: pd.DataFrame) -> pd.DataFrame:
    out = []
    for s, g in latest.groupby("series"):
        per = pd.PeriodIndex(sorted(set(g.period)), freq="M")
        full = pd.period_range(per.min(), per.max(), freq="M")
        out.append({"series": s, "first_period": str(per.min()), "last_period": str(per.max()),
                    "n_months": len(per), "n_gaps": len(full) - len(per),
                    "n_vintages_used": g.source_file.nunique()})
    return pd.DataFrame(out)


def revisions(pv: pd.DataFrame) -> pd.DataFrame:
    """Compare vintages that report the same series and period."""
    pv = pv.assign(_vk=vintage_key(pv)).sort_values(["series", "period", "_vk"])
    out = []
    for s, g in pv.groupby("series"):
        stats = g.groupby("period").value.agg(["count", "min", "max", "first", "last"])
        multi = stats[stats["count"] > 1]
        if multi.empty:
            continue
        spread = multi["max"] - multi["min"]
        scale = multi["last"].abs().where(multi["last"].abs() > 0)
        rel = (multi["last"] - multi["first"]).abs() / scale * 100
        out.append({"series": s, "periods_multi_vintage": len(multi),
                    "periods_revised": int((spread > 1e-6 * multi["max"].abs().clip(lower=1)).sum()),
                    "first_vs_latest_median_pct": round(float(rel.median()), 4) if rel.notna().any() else None,
                    "first_vs_latest_p90_pct": round(float(rel.quantile(0.9)), 4) if rel.notna().any() else None,
                    "periods_first_vs_latest_gt_1pct": int((rel > 1).sum()),
                    "max_abs_spread": float(spread.max())})
    return pd.DataFrame(out)


def main() -> int:
    files = sorted(p for p in TABLES.iterdir() if p.suffix.lower() in (".xls", ".xlsx") and not SKIP_FILES.search(p.name))
    store, report = Store(), []
    for k, p in enumerate(files, 1):
        parse_file(p, store, report)
        r = report[-1]
        print(f"[{k:3d}/{len(files)}] {r['ref'] or '???????'} {p.name[:90]} "
              f"{r['status'] if r['status'] != 'ok' else ''}", flush=True)
    long = pd.DataFrame(store.rows)
    long = long.drop_duplicates(["series", "period", "source_file", "sheet", "_key", "value"])
    # monetary operation components: sum sub instruments (tenors) within a sheet
    is_mo = long.series.str.startswith("mo_")
    mo = long[is_mo].drop_duplicates(["series", "period", "source_file", "sheet", "_key"])
    keep = [c for c in long.columns if c not in ("value", "_key")]
    mo = mo.groupby(keep, as_index=False, dropna=False).value.sum()
    long = pd.concat([long[~is_mo], mo], ignore_index=True)
    long = long.drop_duplicates(["series", "period", "value", "source_file", "sheet"])
    pv = per_vintage(long)
    cols = ["series", "period", "value", "unit", "basis", "source_file", "sheet", "vintage_reference_month"]
    out_long = long[cols].sort_values(["series", "period", "vintage_reference_month", "source_file", "sheet"])
    out_long = out_long.reset_index(drop=True)
    wide, latest = build_wide(pv)
    wide, latest = add_mo_totals(wide, latest)
    wide, latest = add_flows(wide, latest)
    cov = coverage(latest)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_long.to_csv(OUT_LONG, index=False)
    wide.to_csv(OUT_WIDE)
    cov.to_csv(OUT_COV, index=False)

    rep = pd.DataFrame(report)
    print("\nfiles read:", len(rep), " errors:", (rep.status != "ok").sum())
    for _, r in rep[rep.status != "ok"].iterrows():
        print("  ", r.file, r.status)
    no_ref = rep[rep.ref == ""]
    if len(no_ref):
        print("files without a parsed reference month:", list(no_ref.file))
    print("\ncoverage (latest vintage table)")
    print(cov.to_string(index=False))
    print("\nrevisions across vintages (per vintage values)")
    print(revisions(pv).to_string(index=False))
    validate(wide, pv)
    print(f"\nwrote {OUT_LONG.relative_to(ROOT)} ({len(out_long)} rows), {OUT_WIDE.relative_to(ROOT)} "
          f"({len(wide)} periods), {OUT_COV.relative_to(ROOT)}")
    return 0


def validate(wide: pd.DataFrame, pv: pd.DataFrame):
    print("\nvalidation")
    if INTERBANK_REF.exists() and "interbank_rate" in wide:
        ref = pd.read_csv(INTERBANK_REF).set_index("month").interbank_rate
        both = pd.concat([wide.interbank_rate, ref], axis=1, keys=["cme", "qeb"]).dropna()
        d = (both.cme - both.qeb).abs()
        print(f"  interbank_rate vs nepal_interbank_monthly.csv: {len(both)} overlapping months, "
              f"mean abs diff {d.mean():.4f}, months with diff > 0.05: {(d > 0.05).sum()}")
    if "workers_remittances" in wide:
        jul = wide.workers_remittances[wide.index.str.endswith("-07")].dropna()
        print("  workers_remittances fiscal year totals (July cumulative, Rs million):")
        prev = None
        for p, v in jul.items():
            flag = "" if prev is None or v >= prev else "  (decline)"
            print(f"    FY ending {p}: {v:,.0f}{flag}")
            prev = v


if __name__ == "__main__":
    sys.exit(main())
